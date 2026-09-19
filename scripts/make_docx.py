# -*- coding: utf-8 -*-
"""Generate a paper knowledge-point report (.docx) from a JSON spec.
Usage: python make_docx.py <spec.json> <out.docx>
Spec format:
{
  "title_en": "...", "title_cn": "...",
  "journal": "...", "year": "...", "doi": "...",
  "authors": ["..."], "affiliations": ["..."],
  "data_methods": "...(markdown-ish plain text, \\n separated paragraphs)",
  "figures": [{"path": "abs/path.png", "caption": "Fig. 1 ...", "note": "结论/解读"}],
  "figures_note": "optional note when no figures available",
  "conclusions": ["bullet1", "bullet2"]
}
"""
import sys, os, json
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

def set_cn_font(run, size=None, bold=None, color=None):
    run.font.name = "Times New Roman"
    from docx.oxml.ns import qn
    r = run._element.rPr
    rF = r.find(qn("w:rFonts"))
    if rF is None:
        from docx.oxml import OxmlElement
        rF = OxmlElement("w:rFonts")
        r.append(rF)
    rF.set(qn("w:eastAsia"), "微软雅黑")
    if size: run.font.size = Pt(size)
    if bold is not None: run.font.bold = bold
    if color: run.font.color.rgb = RGBColor(*color)

def para(doc, text, size=10.5, bold=False, color=None, align=None, space_after=6):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_cn_font(run, size, bold, color)
    if align: p.alignment = align
    p.paragraph_format.space_after = Pt(space_after)
    return p

def heading(doc, text, size=13, color=(0x1F, 0x4E, 0x79)):
    p = doc.add_paragraph()
    run = p.add_run(text)
    set_cn_font(run, size, True, color)
    p.paragraph_format.space_before = Pt(12)
    p.paragraph_format.space_after = Pt(6)
    return p

def main():
    spec = json.load(open(sys.argv[1], encoding="utf-8"))
    out = sys.argv[2]
    doc = Document()
    # page margins
    for s in doc.sections:
        s.top_margin = Cm(2); s.bottom_margin = Cm(2)
        s.left_margin = Cm(2.2); s.right_margin = Cm(2.2)

    para(doc, spec["title_cn"], 15, True, (0, 0, 0), WD_ALIGN_PARAGRAPH.CENTER, 2)
    para(doc, spec["title_en"], 12, False, (0x44, 0x44, 0x44), WD_ALIGN_PARAGRAPH.CENTER, 4)
    info = f'{spec.get("journal","")} ({spec.get("year","")})  DOI: {spec.get("doi","")}'
    para(doc, info, 9.5, False, (0x77, 0x77, 0x77), WD_ALIGN_PARAGRAPH.CENTER, 10)

    heading(doc, "一、作者与单位")
    para(doc, "作者：" + "；".join(spec["authors"]), 10.5)
    for i, aff in enumerate(spec["affiliations"], 1):
        para(doc, f"{i}. {aff}", 10)

    heading(doc, "二、数据与方法")
    for ptxt in spec["data_methods"].split("\n"):
        ptxt = ptxt.strip()
        if ptxt:
            para(doc, ptxt, 10.5)

    heading(doc, "三、关键图件与解读")
    figs = spec.get("figures", [])
    if not figs:
        para(doc, spec.get("figures_note", "（本篇未能获取原文图件，仅基于文字内容总结。）"), 10, False, (0x88, 0x44, 0x00))
    for fg in figs:
        path = fg["path"]
        if os.path.exists(path):
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run()
            try:
                run.add_picture(path, width=Cm(14.5))
            except Exception:
                run.add_picture(path, width=Cm(12))
        para(doc, fg["caption"], 9.5, True, (0x1F, 0x4E, 0x79), WD_ALIGN_PARAGRAPH.CENTER, 2)
        if fg.get("note"):
            para(doc, "解读与相关结论：" + fg["note"], 10, False, None, None, 10)

    heading(doc, "四、主要结论")
    for i, c in enumerate(spec["conclusions"], 1):
        para(doc, f"{i}. {c}", 10.5)

    doc.save(out)
    print(json.dumps({"ok": True, "out": out}))

if __name__ == "__main__":
    main()
